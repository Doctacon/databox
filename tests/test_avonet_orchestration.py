from __future__ import annotations

import json
import re
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import duckdb
import pytest
from databox.config.sources import by_name
from databox.destinations import (
    dlt_destination,
    dlt_pipeline,
    prepare_dlt_source,
    quack_ingest_session,
)
from databox.orchestration.definitions import defs
from databox.orchestration.domains import avonet
from databox.orchestration.parallel_refresh import execute_parallel_refresh
from databox.quality.platform_health_codegen import render as render_platform_health
from databox_sources.avonet import source
from dlt.pipeline.exceptions import PipelineStepFailed
from openpyxl import Workbook


def _row(name: str, avibase_id: str) -> list[object]:
    return [
        name,
        "Accipitridae",
        "Accipitriformes",
        avibase_id,
        5,
        2,
        0,
        3,
        4,
        27.7,
        17.8,
        10.6,
        14.7,
        62,
        235.2,
        81.8,
        159.5,
        33.9,
        169,
        248.8,
        "Dunning",
        "NA",
        "NO",
        "NA",
        "NA",
        "Forest",
        1,
        2,
        "Carnivore",
        "Vertivore",
        "Insessorial",
    ]


def _workbook(*identities: tuple[str, str]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = source.AVONET_WORKSHEET
    worksheet.append(list(source.EXPECTED_HEADERS))
    for name, avibase_id in identities:
        worksheet.append(_row(name, avibase_id))
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _pipeline(database: Path, pipelines_dir: Path):
    return dlt_pipeline(
        pipeline_name="avonet_test",
        destination=dlt_destination(str(database)),
        dataset_name=avonet._AVONET_STAGING_SCHEMA,
        pipelines_dir=str(pipelines_dir),
    )


def _run_snapshot(
    *,
    database: Path,
    pipeline,
    monkeypatch: pytest.MonkeyPatch,
    identities: tuple[tuple[str, str], ...],
    publish_expected_rows: int | None = None,
) -> None:
    body = _workbook(*identities)

    def download(destination: Path) -> None:
        destination.write_bytes(body)

    expected_rows = len(identities)
    monkeypatch.setattr(source, "download_avonet_workbook", download)
    monkeypatch.setattr(source, "AVONET_EXPECTED_ROWS", expected_rows)
    with avonet.avonet_staged_publish(
        str(database),
        expected_rows=(expected_rows if publish_expected_rows is None else publish_expected_rows),
    ):
        with patch("databox.destinations.quack.settings.quack_shared_server", False):
            with quack_ingest_session(avonet._AVONET_STAGING_SCHEMA, str(database)):
                load = pipeline.run(prepare_dlt_source(source.avonet_source()))
    assert not load.has_failed_jobs


def _business_rows(database: Path) -> list[tuple[str, str]]:
    with duckdb.connect(str(database), read_only=True) as connection:
        return connection.execute(
            """
            SELECT source_scientific_name, avibase_id
            FROM raw_avonet.species_traits
            ORDER BY avibase_id
            """
        ).fetchall()


def _final_snapshot(database: Path) -> dict[str, tuple[tuple[str, ...], tuple[tuple, ...]]]:
    with duckdb.connect(str(database), read_only=True) as connection:
        tables = [
            row[0]
            for row in connection.execute(
                """
                SELECT table_name
                FROM information_schema.tables
                WHERE table_schema = 'raw_avonet' AND table_type = 'BASE TABLE'
                ORDER BY table_name
                """
            ).fetchall()
        ]
        return {
            table: (
                tuple(
                    row[0]
                    for row in connection.execute(
                        """
                        SELECT column_name
                        FROM information_schema.columns
                        WHERE table_schema = 'raw_avonet' AND table_name = ?
                        ORDER BY ordinal_position
                        """,
                        [table],
                    ).fetchall()
                ),
                tuple(
                    connection.execute(f"SELECT * FROM raw_avonet.{table} ORDER BY ALL").fetchall()
                ),
            )
            for table in tables
        }


def _assert_published_layout(database: Path) -> None:
    with duckdb.connect(str(database), read_only=True) as connection:
        tables = {
            row[0]
            for row in connection.execute(
                """
                SELECT table_name
                FROM information_schema.tables
                WHERE table_schema = 'raw_avonet' AND table_type = 'BASE TABLE'
                """
            ).fetchall()
        }
        assert tables == {"species_traits", "_dlt_loads", "_dlt_version"}
        columns = tuple(
            row[0]
            for row in connection.execute(
                """
                SELECT column_name
                FROM information_schema.columns
                WHERE table_schema = 'raw_avonet' AND table_name = 'species_traits'
                ORDER BY ordinal_position
                """
            ).fetchall()
        )
        assert columns == avonet._AVONET_EXPECTED_COLUMNS


def _assert_no_transient_relations(database: Path) -> None:
    with duckdb.connect(str(database), read_only=True) as connection:
        assert connection.execute(
            """
            SELECT count(*)
            FROM information_schema.schemata
            WHERE schema_name = 'raw_avonet_staging'
            """
        ).fetchone() == (0,)
        assert connection.execute(
            """
            SELECT count(*)
            FROM information_schema.tables
            WHERE table_schema = 'raw_avonet_staging'
            """
        ).fetchone() == (0,)
        assert connection.execute(
            """
            SELECT count(*)
            FROM information_schema.tables
            WHERE table_schema = 'main' AND table_name LIKE '\\_dlt%' ESCAPE '\\'
            """
        ).fetchone() == (0,)


def test_avonet_is_independent_unscheduled_and_not_in_parallel_refresh() -> None:
    registered = by_name("avonet")
    assert registered is not None
    assert registered.raw_tables == ("species_traits",)
    assert registered.scheduled is False
    assert registered.parallel_refresh is False
    assert avonet.ingest_job.name == "avonet_ingest"
    assert not hasattr(avonet, "schedule")
    assert not hasattr(avonet, "daily_pipeline")
    assert defs.get_job_def("avonet_ingest").name == "avonet_ingest"
    with pytest.raises(ValueError, match="Unknown sources: avonet"):
        execute_parallel_refresh(["avonet"])
    assert "raw_avonet" not in render_platform_health()


def test_avonet_schema_artifacts_match_normalized_resource_and_annotations() -> None:
    schema_dir = Path(".schema/environmental_observations")
    dbml = (schema_dir / "avonet.dbml").read_text()
    species_table = dbml.split('Table "species_traits"', maxsplit=1)[1]
    columns = re.findall(r'^  "([^"]+)" ', species_table, flags=re.MULTILINE)
    assert columns == [*source._COLUMNS, "_dlt_load_id", "_dlt_id"]
    assert "millimetres" in species_table
    assert "grams" in species_table
    assert "1 dense, 2 semi-open, 3 open" in species_table
    assert "1 sedentary, 2 partial migrant, 3 migratory" in species_table
    taxonomy = json.loads((schema_dir / "taxonomy.json").read_text())
    assert taxonomy["BirdSpeciesTraits"]["natural_key"] == "avibase_id"
    assert taxonomy["BirdSpeciesTraits"]["tables"] == [
        {"table": "species_traits", "source_pipeline": "avonet", "role": "primary"}
    ]
    ontology = (schema_dir / "ontology.md").read_text()
    assert "## BirdSpeciesTraits" in ontology
    assert "global AVONET species averages" in ontology


def test_production_route_replaces_snapshot_and_is_idempotent(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    database = tmp_path / "databox.duckdb"
    pipeline = _pipeline(database, tmp_path / "pipelines")
    first = (("Accipiter albogularis", "AVIBASE-ONE"),)

    _run_snapshot(
        database=database,
        pipeline=pipeline,
        monkeypatch=monkeypatch,
        identities=first,
    )
    _run_snapshot(
        database=database,
        pipeline=pipeline,
        monkeypatch=monkeypatch,
        identities=first,
    )

    assert _business_rows(database) == list(first)
    _assert_published_layout(database)
    _assert_no_transient_relations(database)


def test_production_route_removes_old_rows_instead_of_retaining_append_history(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    database = tmp_path / "databox.duckdb"
    pipeline = _pipeline(database, tmp_path / "pipelines")
    _run_snapshot(
        database=database,
        pipeline=pipeline,
        monkeypatch=monkeypatch,
        identities=(
            ("Accipiter albogularis", "AVIBASE-ONE"),
            ("Buteo jamaicensis", "AVIBASE-TWO"),
        ),
    )
    replacement = (("Corvus corax", "AVIBASE-THREE"),)
    _run_snapshot(
        database=database,
        pipeline=pipeline,
        monkeypatch=monkeypatch,
        identities=replacement,
    )

    assert _business_rows(database) == list(replacement)
    _assert_published_layout(database)
    _assert_no_transient_relations(database)


def test_extraction_failure_preserves_prior_final_and_cleans_staging(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    database = tmp_path / "databox.duckdb"
    pipeline = _pipeline(database, tmp_path / "pipelines")
    _run_snapshot(
        database=database,
        pipeline=pipeline,
        monkeypatch=monkeypatch,
        identities=(("Accipiter albogularis", "AVIBASE-ONE"),),
    )
    before = _final_snapshot(database)

    def fail_download(destination: Path) -> None:
        raise ValueError("fixture extraction failure")

    monkeypatch.setattr(source, "download_avonet_workbook", fail_download)
    with pytest.raises(PipelineStepFailed):
        with avonet.avonet_staged_publish(str(database), expected_rows=1):
            with patch("databox.destinations.quack.settings.quack_shared_server", False):
                with quack_ingest_session(avonet._AVONET_STAGING_SCHEMA, str(database)):
                    pipeline.run(prepare_dlt_source(source.avonet_source()))

    assert _final_snapshot(database) == before
    _assert_no_transient_relations(database)


def test_injected_staged_load_lifecycle_failure_preserves_prior_final(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    database = tmp_path / "databox.duckdb"
    pipeline = _pipeline(database, tmp_path / "pipelines")
    _run_snapshot(
        database=database,
        pipeline=pipeline,
        monkeypatch=monkeypatch,
        identities=(("Accipiter albogularis", "AVIBASE-ONE"),),
    )
    before = _final_snapshot(database)
    body = _workbook(("Corvus corax", "AVIBASE-TWO"))

    def download(destination: Path) -> None:
        destination.write_bytes(body)

    monkeypatch.setattr(source, "download_avonet_workbook", download)
    monkeypatch.setattr(source, "AVONET_EXPECTED_ROWS", 1)
    with pytest.raises(RuntimeError, match="injected staged-load failure"):
        with avonet.avonet_staged_publish(str(database), expected_rows=1):
            with patch("databox.destinations.quack.settings.quack_shared_server", False):
                with quack_ingest_session(avonet._AVONET_STAGING_SCHEMA, str(database)):
                    pipeline.run(prepare_dlt_source(source.avonet_source()))
                    raise RuntimeError("injected staged-load failure")

    assert _final_snapshot(database) == before
    _assert_no_transient_relations(database)


def test_validation_failure_preserves_prior_final_and_first_failure_publishes_nothing(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    database = tmp_path / "databox.duckdb"
    pipeline = _pipeline(database, tmp_path / "pipelines")
    with pytest.raises(RuntimeError, match="row or uniqueness"):
        _run_snapshot(
            database=database,
            pipeline=pipeline,
            monkeypatch=monkeypatch,
            identities=(("Accipiter albogularis", "AVIBASE-ONE"),),
            publish_expected_rows=2,
        )
    with duckdb.connect(str(database), read_only=True) as connection:
        assert not avonet._table_exists(connection, "raw_avonet", "species_traits")
    _assert_no_transient_relations(database)

    _run_snapshot(
        database=database,
        pipeline=pipeline,
        monkeypatch=monkeypatch,
        identities=(("Accipiter albogularis", "AVIBASE-ONE"),),
    )
    before = _final_snapshot(database)
    with pytest.raises(RuntimeError, match="row or uniqueness"):
        _run_snapshot(
            database=database,
            pipeline=pipeline,
            monkeypatch=monkeypatch,
            identities=(("Corvus corax", "AVIBASE-TWO"),),
            publish_expected_rows=2,
        )
    assert _final_snapshot(database) == before
    _assert_no_transient_relations(database)


def test_mid_publish_failure_rolls_back_final_and_cleans_staging(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    database = tmp_path / "databox.duckdb"
    pipeline = _pipeline(database, tmp_path / "pipelines")
    _run_snapshot(
        database=database,
        pipeline=pipeline,
        monkeypatch=monkeypatch,
        identities=(("Accipiter albogularis", "AVIBASE-ONE"),),
    )
    before = _final_snapshot(database)
    original_replace = avonet._replace_avonet_final_table

    def fail_after_business_replace(connection, table: str) -> None:
        original_replace(connection, table)
        if table == avonet._AVONET_BUSINESS_TABLE:
            raise RuntimeError("injected mid-publish failure")

    monkeypatch.setattr(avonet, "_replace_avonet_final_table", fail_after_business_replace)
    with pytest.raises(RuntimeError, match="injected mid-publish failure"):
        _run_snapshot(
            database=database,
            pipeline=pipeline,
            monkeypatch=monkeypatch,
            identities=(("Corvus corax", "AVIBASE-TWO"),),
        )

    assert _final_snapshot(database) == before
    _assert_no_transient_relations(database)


def test_crash_residue_is_cleared_before_next_run(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    database = tmp_path / "databox.duckdb"
    with duckdb.connect(str(database)) as connection:
        connection.execute("CREATE SCHEMA raw_avonet_staging")
        connection.execute("CREATE TABLE raw_avonet_staging.species_traits (bad INTEGER)")
    pipeline = _pipeline(database, tmp_path / "pipelines")

    _run_snapshot(
        database=database,
        pipeline=pipeline,
        monkeypatch=monkeypatch,
        identities=(("Accipiter albogularis", "AVIBASE-ONE"),),
    )

    assert _business_rows(database) == [("Accipiter albogularis", "AVIBASE-ONE")]
    _assert_published_layout(database)
    _assert_no_transient_relations(database)
